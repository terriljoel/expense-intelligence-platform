from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(r"C:\Users\HiWi\Desktop\Terril\01_nextcloud\Germany\DATA SCIENCE\03 PROJECT\02 Expense Tracker\expense-intelligence-platform")
OUT = ROOT / "Implementation_Status_and_Future_Plan.docx"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    tbl_header = OxmlElement("w:tblHeader")
    tbl_header.set(qn("w:val"), "true")
    tr_pr.append(tbl_header)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement("w:tcMar")
    for name, value in [("top", top), ("start", start), ("bottom", bottom), ("end", end)]:
      node = OxmlElement(f"w:{name}")
      node.set(qn("w:w"), str(value))
      node.set(qn("w:type"), "dxa")
      tc_mar.append(node)
    tc_pr.append(tc_mar)


def format_run(run, size=11, bold=False, color="000000"):
    run.font.name = "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "Calibri")
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = RGBColor.from_string(color)


def format_para(p, before=0, after=6, line=1.25, keep=False):
    fmt = p.paragraph_format
    fmt.space_before = Pt(before)
    fmt.space_after = Pt(after)
    fmt.line_spacing = line
    fmt.keep_together = keep
    fmt.keep_with_next = keep


doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.5)
sec.page_height = Inches(11)
for margin in (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin):
    pass
sec.top_margin = Inches(1)
sec.bottom_margin = Inches(1)
sec.left_margin = Inches(1)
sec.right_margin = Inches(1)
sec.header_distance = Inches(0.492)
sec.footer_distance = Inches(0.492)

styles = doc.styles
styles["Normal"].font.name = "Calibri"
styles["Normal"].font.size = Pt(11)

title = doc.add_paragraph()
title.alignment = WD_ALIGN_PARAGRAPH.LEFT
format_para(title, after=3, line=1.0)
tr = title.add_run("Expense Intelligence Platform")
format_run(tr, size=26, bold=False, color="000000")

sub = doc.add_paragraph()
format_para(sub, before=0, after=10, line=1.0)
sr = sub.add_run("Implementation status and future plan")
format_run(sr, size=12, bold=False, color="555555")

body = [
    ("Current status", [
        "Milestone 1 is complete: local PostgreSQL, FastAPI, authentication, households, merchants, payment methods, receipts, receipt items, receipt payments, OCR capture flow, duplicate checks, and the working server-rendered UI are in place.",
        "The app already supports multi-user and multi-household behavior, including admin access and household-scoped data access.",
        "OCR processing is integrated with Azure Document Intelligence and includes image preparation/compression before sending a document to Azure.",
        "A local fallback path exists for image storage, and the codebase is modularized so integrations stay separated from the UI.",
    ]),
    ("What has been implemented", [
        "Database layer: PostgreSQL schema, migrations, and core entities are working.",
        "API layer: FastAPI routes for auth, households, merchants, payment methods, receipts, receipt items, receipt payments, captures, income, and maintenance are implemented.",
        "UI layer: a browser-based interface is working on desktop and mobile.",
        "Operational flow: Docker-based local development, pytest coverage for key flows, and manual database inspection/debugging are already usable.",
    ]),
    ("Design decisions already made", [
        "PostgreSQL is the primary relational store.",
        "Household scoping is enforced across the app.",
        "Merchant normalization is household-aware.",
        "Payment methods are owned by users, while receipt payments can reference them for a shared household receipt.",
        "OCR upload handling now compresses the Azure-bound copy only, so original stored images remain intact.",
    ]),
    ("Next implementation phase", [
        "Finish Milestone 2: normalized product foundation with merchant aliases, product identity, item mapping, package metadata, and price history.",
        "Add deterministic merchant and item resolution so common matches do not require external research.",
        "Add review workflows for unresolved or low-confidence matches.",
        "Extend enrichment for categories, product metadata, and nutrition-related fields.",
        "Build analytics and dashboard views for expenditure breakdowns, top vendors, price trends, and purchase insights.",
    ]),
    ("Later roadmap", [
        "Controlled research agent for unresolved merchants and products.",
        "Evaluation gates for when to use RAG or embeddings.",
        "Export, archive, and purge flows.",
        "Production hardening, budget controls, backups, and observability.",
    ]),
]

for heading, items in body:
    p = doc.add_paragraph()
    format_para(p, before=10, after=3, line=1.0)
    r = p.add_run(heading)
    format_run(r, size=16, bold=False)
    for item in items:
        bp = doc.add_paragraph(style="List Bullet")
        format_para(bp, before=0, after=4, line=1.15)
        br = bp.add_run(item)
        format_run(br, size=11)

doc.add_paragraph()

tbl = doc.add_table(rows=1, cols=2)
tbl.autofit = False
tbl.columns[0].width = Inches(1.8)
tbl.columns[1].width = Inches(4.7)
hdr = tbl.rows[0].cells
hdr[0].text = "Area"
hdr[1].text = "Status"
for c in hdr:
    set_cell_shading(c, "E8EEF5")
    set_cell_margins(c)
    for p in c.paragraphs:
        format_para(p, after=0, line=1.0)
        for r in p.runs:
            format_run(r, size=11, bold=True)
set_repeat_table_header(tbl.rows[0])

rows = [
    ("Database", "Done: PostgreSQL schema, migrations, and local Docker setup."),
    ("API", "Done: FastAPI routes and vertical-slice workflows are in place."),
    ("UI", "Done: working browser UI with mobile access."),
    ("OCR", "Done: Azure OCR integration plus image compression for Azure limits."),
    ("Next", "Build normalized product and resolution layers."),
]
for left, right in rows:
    row = tbl.add_row().cells
    row[0].text = left
    row[1].text = right
    for c in row:
        set_cell_margins(c)
        for p in c.paragraphs:
            format_para(p, after=0, line=1.15)
            for r in p.runs:
                format_run(r, size=11)

for row in tbl.rows:
    for cell in row.cells:
        for tc in cell._tc.iter():
            if tc.tag.endswith("tcBorders"):
                break

doc.save(OUT)
print(OUT)
